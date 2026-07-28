"""
Excel Export

负责：

1. 创建Excel
2. 写入简历
3. 保存xlsx
"""

from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font


BASE = Path(__file__).resolve().parent.parent.parent

OUTPUT_DIR = BASE / "data" / "output"

OUTPUT_DIR.mkdir(
    parents=True,
    exist_ok=True
)


def export_excel(resumes: list[dict]) -> str:
    """
    导出Excel
    """

    workbook = Workbook()

    sheet = workbook.active

    sheet.title = "Resume"

    headers = [
        "ID",
        "姓名",
        "学校",
        "专业",
        "研究方向",
        "技能",
        "论文",
        "创建时间"
    ]

    for col, header in enumerate(headers, start=1):

        cell = sheet.cell(
            row=1,
            column=col
        )

        cell.value = header
        cell.font = Font(bold=True)

    row = 2

    for resume in resumes:

        skills = resume["skills"]

        if isinstance(skills, list):
            skills = "、".join(skills)

        sheet.cell(row=row, column=1).value = resume["id"]
        sheet.cell(row=row, column=2).value = resume["name"]
        sheet.cell(row=row, column=3).value = resume["school"]
        sheet.cell(row=row, column=4).value = resume["major"]
        sheet.cell(row=row, column=5).value = resume["research_direction"]
        sheet.cell(row=row, column=6).value = skills
        sheet.cell(row=row, column=7).value = resume["papers"]
        sheet.cell(row=row, column=8).value = resume["created_at"]

        row += 1

    widths = {
        "A": 8,
        "B": 12,
        "C": 25,
        "D": 25,
        "E": 30,
        "F": 30,
        "G": 10,
        "H": 22,
    }

    for col, width in widths.items():

        sheet.column_dimensions[col].width = width

    filename = (
        "resume_export_"
        + datetime.now().strftime("%Y%m%d_%H%M%S")
        + ".xlsx"
    )

    filepath = OUTPUT_DIR / filename

    workbook.save(filepath)

    return str(filepath)


if __name__ == "__main__":

    demo = [

        {
            "id": 1,
            "name": "张伟",
            "school": "上海交通大学",
            "major": "计算机科学与技术",
            "research_direction": "人工智能",
            "skills": [
                "Python",
                "FastAPI",
                "Docker"
            ],
            "papers": "3篇",
            "created_at": "2026-07-16"
        }

    ]

    print("=" * 60)

    print(export_excel(demo))

    print("=" * 60)