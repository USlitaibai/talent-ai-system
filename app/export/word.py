"""
Word 导出
"""

from docx import Document


def export_word(data: dict, save_path: str):
    """
    导出 Word
    """

    doc = Document()

    doc.add_heading("简历信息", level=1)

    for key, value in data.items():

        if isinstance(value, list):
            value = "、".join(value)

        doc.add_paragraph(f"{key}：{value}")

    doc.save(save_path)

    return save_path


if __name__ == "__main__":

    demo = {
        "姓名": "张伟",
        "学校": "上海交通大学",
        "专业": "计算机科学",
        "技能": [
            "Python",
            "FastAPI",
            "Docker",
        ],
        "论文": "3篇"
    }

    export_word(
        demo,
        "resume.docx"
    )

    print("Word 导出成功")